"""Create a neat Batch-1 weekend attendance + topic tracker workbook."""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import (
    Font, Fill, PatternFill, Alignment, Border, Side, NamedStyle, Protection
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, Reference, PieChart
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.page import PageMargins
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import DataPoint as DP
from copy import copy

OUT = r"E:\ai_data_rag\ai_data_v2\Batch1_Attendance_Topic_Tracker.xlsx"

# Palette
NAVY = "1B3A4B"
TEAL = "2A9D8F"
GOLD = "E9C46A"
SAND = "F4F1DE"
INK = "264653"
WHITE = "FFFFFF"
LIGHT = "F7F9FB"
ROW_ALT = "EEF6F5"
GREEN = "2E7D32"
GREEN_BG = "C8E6C9"
RED = "C62828"
RED_BG = "FFCDD2"
AMBER = "F9A825"
AMBER_BG = "FFF9C4"
SLATE = "546E7A"
HEADER2 = "264653"
LINE = "D0D7DE"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
med_navy = Border(
    left=Side(style="thin", color=NAVY),
    right=Side(style="thin", color=NAVY),
    top=Side(style="thin", color=NAVY),
    bottom=Side(style="thin", color=NAVY),
)

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_gold = PatternFill("solid", fgColor=GOLD)
fill_sand = PatternFill("solid", fgColor=SAND)
fill_light = PatternFill("solid", fgColor=LIGHT)
fill_alt = PatternFill("solid", fgColor=ROW_ALT)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_header2 = PatternFill("solid", fgColor=HEADER2)
fill_green = PatternFill("solid", fgColor=GREEN_BG)
fill_red = PatternFill("solid", fgColor=RED_BG)
fill_amber = PatternFill("solid", fgColor=AMBER_BG)
fill_ink = PatternFill("solid", fgColor=INK)

font_white_b = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_white_title = Font(name="Calibri", size=20, bold=True, color=WHITE)
font_white_sub = Font(name="Calibri", size=12, color="D8EDE9")
font_ink = Font(name="Calibri", size=11, color=INK)
font_ink_b = Font(name="Calibri", size=11, bold=True, color=INK)
font_small = Font(name="Calibri", size=9, color=SLATE)
font_teal_b = Font(name="Calibri", size=12, bold=True, color=TEAL)
font_navy_b = Font(name="Calibri", size=14, bold=True, color=NAVY)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
left_nowrap = Alignment(horizontal="left", vertical="center", wrap_text=False)

# Weekend class dates: Sat + Sun from 22 Aug 2026 through 27 Sep 2026
# (covers completed + a few upcoming sessions)
def weekend_dates(start: date, end: date):
    d = start
    out = []
    while d <= end:
        if d.weekday() in (5, 6):  # Sat, Sun
            out.append(d)
        d += timedelta(days=1)
    return out

SESSIONS = weekend_dates(date(2026, 8, 22), date(2026, 9, 27))
# 22/23 Aug, 29/30 Aug, 5/6 Sep, 12/13 Sep, 19/20 Sep, 26/27 Sep = 12 sessions

TOPICS = [
    # date, week, topic, subtopic, activity, hours, status
    (date(2026, 8, 22), 1, "Introduction to Cloud", "Why Cloud, IaaS / PaaS / SaaS", "Concept + Q&A", 3, "Completed"),
    (date(2026, 8, 23), 1, "Introduction to Cloud", "Public / Private / Hybrid / Multi-cloud", "Concept + comparison table", 3, "Completed"),
    (date(2026, 8, 29), 2, "Azure Fundamentals", "Subscriptions, Resource Groups, Regions", "Portal walkthrough", 3, "Completed"),
    (date(2026, 8, 30), 2, "Azure Storage", "Blob, ADLS Gen2, containers", "Lab: create storage + upload", 3, "Completed"),
    (date(2026, 9, 5), 3, "Azure Data Factory", "ADF overview, linked services, datasets", "Demo", 3, "Completed"),
    (date(2026, 9, 6), 3, "Azure Data Factory", "Copy Activity (file → ADLS / SQL)", "Lab: Copy pipeline", 3, "Completed"),
    (date(2026, 9, 12), 4, "Azure Data Factory", "Pipelines, triggers, parameters", "Lab: scheduled copy", 3, "Planned"),
    (date(2026, 9, 13), 4, "Azure Data Factory", "Mapping Data Flows (intro)", "Demo + practice", 3, "Planned"),
    (date(2026, 9, 19), 5, "Azure Databricks", "Workspace, clusters, notebooks", "Lab: Spark notebook", 3, "Planned"),
    (date(2026, 9, 20), 5, "Azure Databricks", "Read / write ADLS, Delta basics", "Lab: bronze table", 3, "Planned"),
    (date(2026, 9, 26), 6, "Azure Synapse / SQL", "Dedicated vs serverless, load patterns", "Lab: COPY INTO", 3, "Planned"),
    (date(2026, 9, 27), 6, "End-to-end recap", "Ingest → store → transform", "Mini project + doubt clearing", 3, "Planned"),
]

STUDENTS = [f"Member {i:02d}" for i in range(1, 9)]
TUTOR = "Tutor"

wb = Workbook()

def apply_print(ws, landscape=True, fit=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = fit
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.print_title_rows = "1:4"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

def header_banner(ws, last_col, title, subtitle, merge_end=None):
    end = merge_end or last_col
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=end)
    c1, c2, c3 = ws["A1"], ws["A2"], ws["A3"]
    c1.value = title
    c2.value = subtitle
    c3.value = "Batch 1  ·  Weekend batch (Sat–Sun)  ·  8 members  ·  Attendance + topic tracker"
    c1.font = font_white_title
    c2.font = font_white_sub
    c3.font = Font(name="Calibri", size=10, italic=True, color="B8D4CE")
    c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c1.fill = fill_navy
    c2.fill = fill_navy
    c3.fill = fill_navy
    for col in range(1, end + 1):
        for row in (1, 2, 3):
            ws.cell(row, col).fill = fill_navy
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18

# ---------------------------------------------------------------------------
# Sheet: Roster
# ---------------------------------------------------------------------------
ws_r = wb.active
ws_r.title = "01_Roster"

header_banner(
    ws_r, 6,
    "BATCH 1  ·  CLASS ROSTER",
    "Update names once here. Attendance and summary sheets pick them up automatically.",
)
ws_r.merge_cells("A5:B5")
ws_r["A5"] = "BATCH DETAILS"
ws_r["A5"].font = font_white_b
ws_r["A5"].fill = fill_teal
ws_r["B5"].fill = fill_teal
ws_r["A5"].alignment = left

meta = [
    ("Batch name", "Batch 1"),
    ("Mode", "Weekend (Saturday & Sunday)"),
    ("Strength", 8),
    ("Tutor", TUTOR),
    ("Start date", date(2026, 8, 22)),
    ("Venue / platform", "To be filled"),
    ("Timing", "To be filled (e.g. 10:00–13:00)"),
    ("Marking key", "P = Present   A = Absent   L = Late   E = Excused"),
]
for i, (k, v) in enumerate(meta, start=6):
    ws_r.cell(i, 1, k).font = font_ink_b
    ws_r.cell(i, 1).fill = fill_sand
    ws_r.cell(i, 1).border = thin
    ws_r.cell(i, 1).alignment = left
    cell = ws_r.cell(i, 2, v)
    cell.font = font_ink
    cell.border = thin
    cell.alignment = left
    if isinstance(v, date):
        cell.number_format = "DD-MMM-YYYY"

ws_r.merge_cells("D5:F5")
ws_r["D5"] = "MEMBERS  (edit names in column E)"
ws_r["D5"].font = font_white_b
ws_r["D5"].fill = fill_teal
ws_r["E5"].fill = fill_teal
ws_r["F5"].fill = fill_teal

for col, h, w in [(4, "Sr", 8), (5, "Full name", 28), (6, "Contact / notes", 28)]:
    c = ws_r.cell(6, col, h)
    c.font = font_white_b
    c.fill = fill_header2
    c.alignment = center
    c.border = thin
    ws_r.column_dimensions[get_column_letter(col)].width = w

for i, name in enumerate(STUDENTS, start=1):
    r = 6 + i
    ws_r.cell(r, 4, i).alignment = center
    ws_r.cell(r, 4).font = font_ink_b
    ws_r.cell(r, 5, name).font = font_ink
    ws_r.cell(r, 6, "").font = font_ink
    fill = fill_white if i % 2 else fill_alt
    for col in range(4, 7):
        ws_r.cell(r, col).fill = fill
        ws_r.cell(r, col).border = thin
        ws_r.cell(r, col).alignment = center if col == 4 else left

ws_r.column_dimensions["A"].width = 22
ws_r.column_dimensions["B"].width = 42
ws_r.column_dimensions["C"].width = 4
ws_r.row_dimensions[6].height = 22
note = ws_r["A15"]
ws_r.merge_cells("A15:B16")
ws_r["A15"] = (
    "Tip: Type actual student names in E7:E14. Do not insert/delete rows in that block — "
    "other sheets are linked to these cells."
)
ws_r["A15"].font = font_small
ws_r["A15"].alignment = Alignment(wrap_text=True, vertical="top")
apply_print(ws_r, landscape=True)

# ---------------------------------------------------------------------------
# Sheet: Session plan
# ---------------------------------------------------------------------------
ws_p = wb.create_sheet("02_Session_Plan")
header_banner(
    ws_p, 9,
    "BATCH 1  ·  SESSION PLAN",
    "Topic-wise weekend calendar. Update status after each class.",
)

headers = ["Sr", "Week", "Date", "Day", "Topic", "Sub-topic", "Activity / lab", "Hrs", "Status"]
widths = [6, 8, 14, 12, 26, 42, 32, 8, 14]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws_p.cell(5, i, h)
    c.font = font_white_b
    c.fill = fill_teal
    c.alignment = center
    c.border = thin
    ws_p.column_dimensions[get_column_letter(i)].width = w
ws_p.row_dimensions[5].height = 24

status_dv = DataValidation(type="list", formula1='"Planned,Completed,Rescheduled,Cancelled"', allow_blank=False)
status_dv.error = "Choose a status from the list"
status_dv.errorTitle = "Invalid status"
ws_p.add_data_validation(status_dv)

for i, row in enumerate(TOPICS, start=1):
    r = 5 + i
    dt, week, topic, sub, act, hrs, status = row
    vals = [i, week, dt, dt.strftime("%A"), topic, sub, act, hrs, status]
    for col, v in enumerate(vals, start=1):
        cell = ws_p.cell(r, col, v)
        cell.font = font_ink
        cell.border = thin
        cell.alignment = center if col in (1, 2, 3, 4, 8, 9) else left
        cell.fill = fill_white if i % 2 else fill_alt
    ws_p.cell(r, 3).number_format = "DD-MMM-YYYY"
    status_dv.add(ws_p.cell(r, 9))
    st = ws_p.cell(r, 9)
    if status == "Completed":
        st.fill = fill_green
        st.font = Font(name="Calibri", size=11, bold=True, color=GREEN)
    else:
        st.fill = fill_amber
        st.font = Font(name="Calibri", size=11, bold=True, color="F57F17")
    ws_p.row_dimensions[r].height = 22

# totals
tr = 6 + len(TOPICS)
ws_p.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=7)
ws_p.cell(tr, 1, "Total scheduled hours").font = font_white_b
ws_p.cell(tr, 1).fill = fill_navy
ws_p.cell(tr, 1).alignment = Alignment(horizontal="right", vertical="center", indent=1)
for col in range(1, 8):
    ws_p.cell(tr, col).fill = fill_navy
    ws_p.cell(tr, col).border = thin
ws_p.cell(tr, 8, f"=SUM(H6:H{tr-1})")
ws_p.cell(tr, 8).font = font_white_b
ws_p.cell(tr, 8).fill = fill_navy
ws_p.cell(tr, 8).alignment = center
ws_p.cell(tr, 8).border = thin
ws_p.cell(tr, 9).fill = fill_navy
ws_p.cell(tr, 9).border = thin

ws_p.freeze_panes = "A6"
apply_print(ws_p)

# ---------------------------------------------------------------------------
# Sheet: Attendance
# Students in rows, dates in columns — the usable daily register
# ---------------------------------------------------------------------------
ws_a = wb.create_sheet("03_Attendance")
n_dates = len(SESSIONS)
last_col = 3 + n_dates + 3  # sr, name, role + dates + P + A + %
header_banner(
    ws_a, last_col,
    "BATCH 1  ·  DAILY ATTENDANCE",
    "Mark P / A / L / E under each weekend date. Names come from the Roster sheet.",
)

# Row 5: week bands, Row 6: topic short, Row 7: date headers
# Columns: A Sr, B Name, C then dates, then Present, Absent, %

ws_a.cell(7, 1, "Sr").fill = fill_teal
ws_a.cell(7, 2, "Name").fill = fill_teal
for col in (1, 2):
    ws_a.cell(7, col).font = font_white_b
    ws_a.cell(7, col).alignment = center
    ws_a.cell(7, col).border = thin
    ws_a.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col)
    for rr in (5, 6):
        ws_a.cell(rr, col).fill = fill_navy
        ws_a.cell(rr, col).border = thin

# week grouping colors
week_fills = [
    PatternFill("solid", fgColor="1B3A4B"),
    PatternFill("solid", fgColor="1F4E5F"),
    PatternFill("solid", fgColor="216B6B"),
    PatternFill("solid", fgColor="2A9D8F"),
    PatternFill("solid", fgColor="3AAFA9"),
    PatternFill("solid", fgColor="2C7A7B"),
]

for i, dt in enumerate(SESSIONS):
    col = 3 + i
    topic = TOPICS[i][2]
    sub = TOPICS[i][3]
    week = TOPICS[i][1]
    # week band
    ws_a.cell(5, col, f"W{week}")
    ws_a.cell(5, col).font = Font(name="Calibri", size=9, bold=True, color=WHITE)
    ws_a.cell(5, col).fill = week_fills[(week - 1) % 6]
    ws_a.cell(5, col).alignment = center
    ws_a.cell(5, col).border = thin
    # topic
    ws_a.cell(6, col, topic)
    ws_a.cell(6, col).font = Font(name="Calibri", size=8, bold=True, color=NAVY)
    ws_a.cell(6, col).fill = fill_sand
    ws_a.cell(6, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, textRotation=0)
    ws_a.cell(6, col).border = thin
    # date + weekday
    ws_a.cell(7, col, dt)
    ws_a.cell(7, col).number_format = "DD-MMM\n(DDD)"
    ws_a.cell(7, col).font = font_white_b
    ws_a.cell(7, col).fill = fill_teal
    ws_a.cell(7, col).alignment = center
    ws_a.cell(7, col).border = thin
    ws_a.column_dimensions[get_column_letter(col)].width = 12

# merge week headers
start = 3
for i, dt in enumerate(SESSIONS):
    week = TOPICS[i][1]
    nxt_week = TOPICS[i + 1][1] if i + 1 < len(SESSIONS) else None
    if nxt_week != week:
        end = 3 + i
        if end > start:
            ws_a.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        ws_a.cell(5, start, f"Week {week}")
        start = end + 1

sum_headers = ["Present", "Absent", "%"]
for j, h in enumerate(sum_headers):
    col = 3 + n_dates + j
    ws_a.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col)
    for rr in (5, 6):
        ws_a.cell(rr, col).fill = fill_navy
        ws_a.cell(rr, col).border = thin
    c = ws_a.cell(7, col, h)
    c.font = font_white_b
    c.fill = fill_header2
    c.alignment = center
    c.border = thin
    ws_a.column_dimensions[get_column_letter(col)].width = 11

ws_a.column_dimensions["A"].width = 6
ws_a.column_dimensions["B"].width = 22
ws_a.row_dimensions[5].height = 20
ws_a.row_dimensions[6].height = 32
ws_a.row_dimensions[7].height = 32

# data validation
mark_dv = DataValidation(type="list", formula1='"P,A,L,E"', allow_blank=True)
mark_dv.prompt = "P Present · A Absent · L Late · E Excused"
mark_dv.promptTitle = "Mark attendance"
ws_a.add_data_validation(mark_dv)

first_data = 8
last_data = 7 + 8
date_first_col = 3
date_last_col = 2 + n_dates
p_col = date_last_col + 1
a_col = date_last_col + 2
pct_col = date_last_col + 3

for i in range(8):
    r = first_data + i
    ws_a.cell(r, 1, i + 1).alignment = center
    ws_a.cell(r, 1).font = font_ink_b
    # linked name
    ws_a.cell(r, 2, f"=RosterName")  # placeholder replaced
    ws_a.cell(r, 2, f"='01_Roster'!E{7+i}")
    ws_a.cell(r, 2).alignment = left
    ws_a.cell(r, 2).font = font_ink
    fill = fill_white if i % 2 == 0 else fill_alt
    ws_a.cell(r, 1).fill = fill
    ws_a.cell(r, 2).fill = fill
    ws_a.cell(r, 1).border = thin
    ws_a.cell(r, 2).border = thin
    for col in range(date_first_col, date_last_col + 1):
        cell = ws_a.cell(r, col, "")
        cell.alignment = center
        cell.font = Font(name="Calibri", size=12, bold=True, color=INK)
        cell.border = thin
        cell.fill = fill
        mark_dv.add(cell)
    # Present counts P and L (late still present)
    rng = f"{get_column_letter(date_first_col)}{r}:{get_column_letter(date_last_col)}{r}"
    ws_a.cell(r, p_col, f'=COUNTIF({rng},"P")+COUNTIF({rng},"L")')
    ws_a.cell(r, a_col, f'=COUNTIF({rng},"A")')
    ws_a.cell(r, pct_col, f'=IF(({p_col-1})=0,"",({get_column_letter(p_col)}{r}+COUNTIF({rng},"E"))/{n_dates})')
    # fix pct: sessions held could be counted later; for now / n_dates of marked? Better:
    # % = (P+L+E) / (P+L+A+E) so empty future dates don't drag % down
    ws_a.cell(
        r, pct_col,
        f'=LET(m,COUNTA({rng}),IF(m=0,"",(COUNTIF({rng},"P")+COUNTIF({rng},"L")+COUNTIF({rng},"E"))/m))',
    )
    # LET might fail on older Excel - use simpler:
    ws_a.cell(
        r, pct_col,
        f'=IF((COUNTA({rng}))=0,"",(COUNTIF({rng},"P")+COUNTIF({rng},"L")+COUNTIF({rng},"E"))/COUNTA({rng}))',
    )
    for col in (p_col, a_col, pct_col):
        ws_a.cell(r, col).alignment = center
        ws_a.cell(r, col).border = thin
        ws_a.cell(r, col).fill = fill
        ws_a.cell(r, col).font = font_ink_b
    ws_a.cell(r, pct_col).number_format = "0%"
    ws_a.row_dimensions[r].height = 22

# conditional formatting for marks
p_range = f"{get_column_letter(date_first_col)}{first_data}:{get_column_letter(date_last_col)}{last_data}"
ws_a.conditional_formatting.add(
    p_range,
    CellIsRule(operator="equal", formula=['"P"'], fill=fill_green,
               font=Font(name="Calibri", size=12, bold=True, color=GREEN)),
)
ws_a.conditional_formatting.add(
    p_range,
    CellIsRule(operator="equal", formula=['"A"'], fill=fill_red,
               font=Font(name="Calibri", size=12, bold=True, color=RED)),
)
ws_a.conditional_formatting.add(
    p_range,
    CellIsRule(operator="equal", formula=['"L"'], fill=fill_amber,
               font=Font(name="Calibri", size=12, bold=True, color="F57F17")),
)
ws_a.conditional_formatting.add(
    p_range,
    CellIsRule(operator="equal", formula=['"E"'],
               fill=PatternFill("solid", fgColor="BBDEFB"),
               font=Font(name="Calibri", size=12, bold=True, color="1565C0")),
)

pct_range = f"{get_column_letter(pct_col)}{first_data}:{get_column_letter(pct_col)}{last_data}"
ws_a.conditional_formatting.add(
    pct_range,
    ColorScaleRule(start_type="num", start_value=0.5, start_color="FFCDD2",
                   mid_type="num", mid_value=0.8, mid_color="FFF9C4",
                   end_type="num", end_value=1, end_color="C8E6C9"),
)

# session totals row
tot = last_data + 1
ws_a.cell(tot, 1, "").fill = fill_navy
ws_a.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=2)
ws_a.cell(tot, 1, "Present count (P+L)")
ws_a.cell(tot, 1).font = font_white_b
ws_a.cell(tot, 1).fill = fill_navy
ws_a.cell(tot, 1).alignment = Alignment(horizontal="right", vertical="center", indent=1)
ws_a.cell(tot, 2).fill = fill_navy
ws_a.cell(tot, 2).border = thin
ws_a.cell(tot, 1).border = thin
for col in range(date_first_col, date_last_col + 1):
    letter = get_column_letter(col)
    ws_a.cell(tot, col, f'=COUNTIF({letter}{first_data}:{letter}{last_data},"P")+COUNTIF({letter}{first_data}:{letter}{last_data},"L")')
    ws_a.cell(tot, col).font = font_white_b
    ws_a.cell(tot, col).fill = fill_navy
    ws_a.cell(tot, col).alignment = center
    ws_a.cell(tot, col).border = thin
for col in range(p_col, pct_col + 1):
    ws_a.cell(tot, col).fill = fill_navy
    ws_a.cell(tot, col).border = thin

# legend
leg = tot + 2
ws_a.merge_cells(start_row=leg, start_column=1, end_row=leg, end_column=6)
ws_a.cell(leg, 1, "Legend:  P = Present (green)    A = Absent (red)    L = Late (amber, counted present)    E = Excused (blue, not treated as absent)    % = (P+L+E) ÷ marked sessions — blank future dates are ignored")
ws_a.cell(leg, 1).font = font_small
ws_a.cell(leg, 1).alignment = left

ws_a.freeze_panes = "C8"
ws_a.auto_filter.ref = f"A7:{get_column_letter(pct_col)}{last_data}"
apply_print(ws_a)
ws_a.page_setup.fitToHeight = 1

# ---------------------------------------------------------------------------
# Sheet: Topic x Member matrix (what they tried in screenshot, done cleanly)
# ---------------------------------------------------------------------------
ws_t = wb.create_sheet("04_Topic_Matrix")
# Layout: left = topic info, then 8 member columns with one mark per session row
last_c = 6 + 8
header_banner(
    ws_t, last_c,
    "BATCH 1  ·  TOPIC-WISE ATTENDANCE MATRIX",
    "One row per class. Member columns are linked to the daily attendance sheet — do not type here.",
)

hrow = 5
heads = ["Sr", "Date", "Day", "Topic", "Sub-topic", "Activity"]
for i, h in enumerate(heads, 1):
    c = ws_t.cell(hrow, i, h)
    c.font = font_white_b
    c.fill = fill_teal
    c.alignment = center
    c.border = thin

for i in range(8):
    col = 7 + i
    # name from roster
    c = ws_t.cell(hrow, col, f"='01_Roster'!E{7+i}")
    c.font = font_white_b
    c.fill = fill_header2
    c.alignment = center
    c.border = thin
    ws_t.column_dimensions[get_column_letter(col)].width = 12

widths_t = [6, 13, 11, 24, 40, 28]
for i, w in enumerate(widths_t, 1):
    ws_t.column_dimensions[get_column_letter(i)].width = w
ws_t.row_dimensions[hrow].height = 28

for i, row in enumerate(TOPICS):
    r = 6 + i
    dt, week, topic, sub, act, hrs, status = row
    vals = [i + 1, dt, dt.strftime("%A"), topic, sub, act]
    fill = fill_white if i % 2 == 0 else fill_alt
    for col, v in enumerate(vals, 1):
        cell = ws_t.cell(r, col, v)
        cell.font = font_ink
        cell.border = thin
        cell.fill = fill
        cell.alignment = center if col <= 3 else left
    ws_t.cell(r, 2).number_format = "DD-MMM-YYYY"
    # pull attendance from sheet 03: column 3+i is that date, rows 8-15 students
    att_col = get_column_letter(3 + i)
    for s in range(8):
        col = 7 + s
        cell = ws_t.cell(r, col, f"='03_Attendance'!{att_col}{8+s}")
        cell.alignment = center
        cell.border = thin
        cell.fill = fill
        cell.font = Font(name="Calibri", size=12, bold=True)
    ws_t.row_dimensions[r].height = 22

tm_range = f"G6:N{5+len(TOPICS)}"
ws_t.conditional_formatting.add(
    tm_range,
    CellIsRule(operator="equal", formula=['"P"'], fill=fill_green,
               font=Font(name="Calibri", size=12, bold=True, color=GREEN)),
)
ws_t.conditional_formatting.add(
    tm_range,
    CellIsRule(operator="equal", formula=['"A"'], fill=fill_red,
               font=Font(name="Calibri", size=12, bold=True, color=RED)),
)
ws_t.conditional_formatting.add(
    tm_range,
    CellIsRule(operator="equal", formula=['"L"'], fill=fill_amber,
               font=Font(name="Calibri", size=12, bold=True, color="F57F17")),
)
ws_t.conditional_formatting.add(
    tm_range,
    CellIsRule(operator="equal", formula=['"E"'],
               fill=PatternFill("solid", fgColor="BBDEFB"),
               font=Font(name="Calibri", size=12, bold=True, color="1565C0")),
)

ws_t.freeze_panes = "G6"
apply_print(ws_t)

# ---------------------------------------------------------------------------
# Sheet: Summary
# ---------------------------------------------------------------------------
ws_s = wb.create_sheet("05_Summary")
header_banner(
    ws_s, 8,
    "BATCH 1  ·  ATTENDANCE SUMMARY",
    "Auto-calculated from daily marks. Refresh after you fill the Attendance sheet.",
)

# KPI cards
kpis = [
    (5, "Batch strength", "='01_Roster'!B8"),
    (7, "Sessions planned", str(len(SESSIONS))),
    (9, "Hours planned", "='02_Session_Plan'!H18"),
]
# H18 is total row: 6+12=18 yes

ws_s["A5"] = "KPI"
ws_s["B5"] = "Value"
ws_s["A5"].font = font_white_b
ws_s["B5"].font = font_white_b
ws_s["A5"].fill = fill_teal
ws_s["B5"].fill = fill_teal
ws_s["A5"].alignment = center
ws_s["B5"].alignment = center
ws_s["A5"].border = thin
ws_s["B5"].border = thin

kpis_rows = [
    ("Batch name", "='01_Roster'!B6"),
    ("Members", 8),
    ("Weekend sessions", len(SESSIONS)),
    ("Planned hours", f"='02_Session_Plan'!H{6+len(TOPICS)}"),
    ("Class avg attendance", f"=IFERROR(AVERAGE('03_Attendance'!{get_column_letter(pct_col)}{first_data}:{get_column_letter(pct_col)}{last_data}),\"\")"),
    ("Below 75% (count)", f'=COUNTIF(\'03_Attendance\'!{get_column_letter(pct_col)}{first_data}:{get_column_letter(pct_col)}{last_data},"<0.75")'),
]
for i, (k, v) in enumerate(kpis_rows, start=6):
    ws_s.cell(i, 1, k).font = font_ink_b
    ws_s.cell(i, 1).fill = fill_sand
    ws_s.cell(i, 1).border = thin
    cell = ws_s.cell(i, 2, v)
    cell.font = font_ink_b
    cell.border = thin
    cell.alignment = center
    cell.fill = fill_light
    if i == 10:
        cell.number_format = "0%"

ws_s.column_dimensions["A"].width = 24
ws_s.column_dimensions["B"].width = 22
ws_s.column_dimensions["C"].width = 3

ws_s.merge_cells("D5:G5")
ws_s["D5"] = "MEMBER STANDING"
ws_s["D5"].font = font_white_b
ws_s["D5"].fill = fill_teal
for col in range(4, 8):
    ws_s.cell(5, col).fill = fill_teal
    ws_s.cell(5, col).border = thin

for col, h, w in [(4, "Sr", 6), (5, "Name", 22), (6, "Attendance %", 14), (7, "Standing", 16)]:
    c = ws_s.cell(6, col, h)
    c.font = font_white_b
    c.fill = fill_header2
    c.alignment = center
    c.border = thin
    ws_s.column_dimensions[get_column_letter(col)].width = w

for i in range(8):
    r = 7 + i
    fill = fill_white if i % 2 == 0 else fill_alt
    ws_s.cell(r, 4, i + 1).alignment = center
    ws_s.cell(r, 5, f"='01_Roster'!E{7+i}")
    ws_s.cell(r, 6, f"='03_Attendance'!{get_column_letter(pct_col)}{8+i}")
    ws_s.cell(r, 6).number_format = "0%"
    ws_s.cell(r, 7, f'=IF(F{r}="","",IF(F{r}>=0.9,"Excellent",IF(F{r}>=0.75,"Good",IF(F{r}>=0.6,"Watch","At risk"))))')
    for col in range(4, 8):
        ws_s.cell(r, col).border = thin
        ws_s.cell(r, col).fill = fill
        ws_s.cell(r, col).font = font_ink
        ws_s.cell(r, col).alignment = center if col != 5 else left
    ws_s.cell(r, 5).font = font_ink

stand_range = f"G7:G14"
ws_s.conditional_formatting.add(stand_range, CellIsRule(operator="equal", formula=['"Excellent"'], fill=fill_green, font=Font(bold=True, color=GREEN, name="Calibri", size=11)))
ws_s.conditional_formatting.add(stand_range, CellIsRule(operator="equal", formula=['"Good"'], fill=PatternFill("solid", fgColor="DCEDC8"), font=Font(bold=True, color="558B2F", name="Calibri", size=11)))
ws_s.conditional_formatting.add(stand_range, CellIsRule(operator="equal", formula=['"Watch"'], fill=fill_amber, font=Font(bold=True, color="F57F17", name="Calibri", size=11)))
ws_s.conditional_formatting.add(stand_range, CellIsRule(operator="equal", formula=['"At risk"'], fill=fill_red, font=Font(bold=True, color=RED, name="Calibri", size=11)))

ws_s.conditional_formatting.add(
    "F7:F14",
    ColorScaleRule(start_type="num", start_value=0.5, start_color="FFCDD2",
                   mid_type="num", mid_value=0.8, mid_color="FFF9C4",
                   end_type="num", end_value=1, end_color="C8E6C9"),
)

chart = BarChart()
chart.type = "col"
chart.title = "Attendance % by member"
chart.y_axis.title = None
chart.x_axis.title = None
chart.style = 10
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 1
chart.y_axis.numFmt = "0%"
data = Reference(ws_s, min_col=6, min_row=6, max_row=14)
cats = Reference(ws_s, min_col=5, min_row=7, max_row=14)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.shape = 4
chart.legend = None
chart.width = 15
chart.height = 8
ws_s.add_chart(chart, "A16")

ws_s.merge_cells("D16:G18")
ws_s["D16"] = (
    "How to use\n"
    "1. Roster — type the 8 names.\n"
    "2. Session Plan — adjust topics/tutor if needed; set Status after class.\n"
    "3. Attendance — each weekend, drop-down P / A / L / E. That is the only sheet you mark.\n"
    "4. Topic Matrix and Summary update by formula. Print Attendance or Topic Matrix in landscape A4."
)
ws_s["D16"].alignment = Alignment(wrap_text=True, vertical="top")
ws_s["D16"].font = Font(name="Calibri", size=10, color=INK)

apply_print(ws_s, landscape=True)

# ---------------------------------------------------------------------------
# How to use sheet (short)
# ---------------------------------------------------------------------------
ws_h = wb.create_sheet("00_How_to_use", 0)
header_banner(
    ws_h, 6,
    "BATCH 1  ·  HOW TO USE THIS WORKBOOK",
    "Keep this file as the single register for weekend Batch 1.",
)
ws_h.column_dimensions["A"].width = 4
ws_h.column_dimensions["B"].width = 22
ws_h.column_dimensions["C"].width = 70
ws_h.column_dimensions["D"].width = 18

ws_h["B5"] = "Sheet"
ws_h["C5"] = "What it is for"
ws_h["D5"] = "You type here?"
for col in ("B", "C", "D"):
    ws_h[f"{col}5"].font = font_white_b
    ws_h[f"{col}5"].fill = fill_teal
    ws_h[f"{col}5"].alignment = center
    ws_h[f"{col}5"].border = thin

guides = [
    ("01_Roster", "Batch info and 8 member names. Change names only in column E.", "Yes — names & meta"),
    ("02_Session_Plan", "Weekend calendar: topic, sub-topic, lab/activity, hours, status.", "Yes — topics/status"),
    ("03_Attendance", "Daily register. Members down the side, dates across. Use the drop-down.", "Yes — P / A / L / E"),
    ("04_Topic_Matrix", "Same marks, pivoted as topic rows × member columns (your original layout, cleaned).", "No — linked"),
    ("05_Summary", "Attendance % , standing (Excellent / Good / Watch / At risk), and a chart.", "No — linked"),
]
for i, (a, b, c) in enumerate(guides, start=6):
    ws_h.cell(i, 2, a).font = font_ink_b
    ws_h.cell(i, 3, b).font = font_ink
    ws_h.cell(i, 4, c).font = font_ink
    fill = fill_white if i % 2 == 0 else fill_alt
    for col in range(2, 5):
        ws_h.cell(i, col).fill = fill
        ws_h.cell(i, col).border = thin
        ws_h.cell(i, col).alignment = left
    ws_h.row_dimensions[i].height = 28

ws_h.merge_cells("B12:D14")
ws_h["B12"] = (
    "Marking: P Present · A Absent · L Late (still counts as present) · E Excused (not absent).\n"
    "Dates included: every Saturday & Sunday from 22-Aug-2026 to 27-Sep-2026 (12 sessions / 6 weekends).\n"
    "Topics follow a Cloud → Azure → ADF → Databricks → Synapse path, matching your sample (Cloud types, ADF Copy). "
    "Replace Member 01–08 and tutor name on the Roster. Add more date columns only if you also extend the Session Plan."
)
ws_h["B12"].alignment = Alignment(wrap_text=True, vertical="top")
ws_h["B12"].font = Font(name="Calibri", size=11, color=INK)
ws_h.row_dimensions[12].height = 20
apply_print(ws_h, landscape=True)

# freeze / tab colors
ws_h.sheet_properties.tabColor = NAVY
ws_r.sheet_properties.tabColor = TEAL
ws_p.sheet_properties.tabColor = GOLD
ws_a.sheet_properties.tabColor = "2E7D32"
ws_t.sheet_properties.tabColor = "1565C0"
ws_s.sheet_properties.tabColor = "6D4C41"

wb.save(OUT)
print("Wrote", OUT)
print("Sessions", len(SESSIONS), [d.isoformat() for d in SESSIONS])
