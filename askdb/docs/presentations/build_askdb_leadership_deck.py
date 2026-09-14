"""Generate Ask DB leadership PowerPoint (light theme).

Run:
  python askdb/docs/presentations/build_askdb_leadership_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

# --- Brand / light theme ----------------------------------------------------
BG = RGBColor(0xF7, 0xF9, 0xFC)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0F, 0x1C, 0x2E)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
ASK = RGBColor(0x1D, 0x6F, 0xE8)  # blue / cyan family
DB = RGBColor(0x0D, 0x9F, 0x8F)  # teal / green family
ACCENT = RGBColor(0x2B, 0x6C, 0xBE)
WARN = RGBColor(0xC2, 0x5E, 0x00)
OK = RGBColor(0x1B, 0x8A, 0x5A)
PURPLE = RGBColor(0x6D, 0x5C, 0xF0)
ORANGE = RGBColor(0xE0, 0x7A, 0x2F)

FOOTER_L = "Ask DB · Capgemini AI Data Platform"
FOOTER_R = "Company Confidential © Capgemini 2026. All rights reserved."

OUT = Path(__file__).resolve().parent / "AskDB_Leadership_Deck_Light.pptx"


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    # Softer corners
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_textbox(slide, left, top, width, height, text, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align
    return box


def add_runs(paragraph, runs: list[tuple[str, RGBColor, bool, int]]) -> None:
    paragraph.clear()
    first = True
    for text, color, bold, size in runs:
        run = paragraph.add_run() if not first or paragraph.runs else paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        if first and paragraph.runs:
            run = paragraph.runs[0]
        else:
            run = paragraph.add_run()
        run.text = text
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        first = False


def brand_title(slide, left, top, width, height, prefix="AI Data Platform — ", size=36):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    # Clear default and build runs
    p.text = ""
    r1 = p.add_run()
    r1.text = prefix
    r1.font.size = Pt(size)
    r1.font.bold = True
    r1.font.color.rgb = INK
    r1.font.name = "Calibri"
    r2 = p.add_run()
    r2.text = "Ask"
    r2.font.size = Pt(size)
    r2.font.bold = True
    r2.font.color.rgb = ASK
    r2.font.name = "Calibri"
    r3 = p.add_run()
    r3.text = " DB"
    r3.font.size = Pt(size)
    r3.font.bold = True
    r3.font.color.rgb = DB
    r3.font.name = "Calibri"
    return box


def footer(slide, page: int, total: int = 16) -> None:
    add_textbox(slide, Inches(0.45), Inches(7.05), Inches(5.5), Inches(0.3), FOOTER_L, size=10, color=MUTED)
    add_textbox(
        slide,
        Inches(5.8),
        Inches(7.05),
        Inches(3.7),
        Inches(0.3),
        f"{FOOTER_R}  {page}",
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def section_title(slide, text: str) -> None:
    add_textbox(slide, Inches(0.5), Inches(0.28), Inches(9), Inches(0.55), text, size=28, bold=True, color=ASK)


def bullet_card(slide, left, top, width, height, title: str, lines: list[str], accent=ASK):
    add_rect(slide, left, top, width, height, CARD, LINE)
    add_rect(slide, left, top, Inches(0.08), height, accent)
    add_textbox(slide, left + Inches(0.22), top + Inches(0.12), width - Inches(0.35), Inches(0.32), title, size=13, bold=True, color=INK)
    body = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.42), width - Inches(0.35), height - Inches(0.55))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.size = Pt(11)
        p.font.color.rgb = MUTED
        p.font.name = "Calibri"
        p.space_after = Pt(4)


def feature_row(slide, left, top, width, title: str, desc: str, color=DB):
    add_rect(slide, left, top, Inches(0.28), Inches(0.28), color)
    add_textbox(slide, left + Inches(0.38), top - Inches(0.02), width - Inches(0.4), Inches(0.26), title, size=12, bold=True, color=INK)
    add_textbox(slide, left + Inches(0.38), top + Inches(0.24), width - Inches(0.4), Inches(0.55), desc, size=10, color=MUTED)


def flow_box(slide, left, top, width, height, text: str, fill=CARD, ink=INK, size=11):
    shape = add_rect(slide, left, top, width, height, fill, LINE)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].text = text
    tf.paragraphs[0].font.size = Pt(size)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ink
    tf.paragraphs[0].font.name = "Calibri"
    shape.text_frame.auto_size = None
    try:
        tf._txBodyBodyPr = None
        shape.text_frame.paragraphs[0].space_before = Pt(6)
    except Exception:
        pass
    return shape


def arrow_right(slide, left, top):
    add_textbox(slide, left, top, Inches(0.3), Inches(0.3), "→", size=16, bold=True, color=ASK, align=PP_ALIGN.CENTER)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = 16

    # 1 Title
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), BG)
    add_rect(s, Inches(0), Inches(0), Inches(0.18), Inches(7.5), ASK)
    add_rect(s, Inches(0.18), Inches(0), Inches(0.08), Inches(7.5), DB)
    brand_title(s, Inches(0.8), Inches(2.2), Inches(11), Inches(0.8), size=40)
    add_textbox(
        s,
        Inches(0.8),
        Inches(3.1),
        Inches(10),
        Inches(0.6),
        "From natural-language questions to governed, actionable business insights",
        size=18,
        color=MUTED,
    )
    add_textbox(
        s,
        Inches(0.8),
        Inches(4.0),
        Inches(10),
        Inches(0.8),
        "Enterprise NLQ over PostgreSQL · Semantic Atlas · Data Trust · Executive Intelligence",
        size=14,
        color=INK,
    )
    add_textbox(s, Inches(0.8), Inches(5.6), Inches(6), Inches(0.3), "Leadership briefing · Capgemini 2026", size=12, color=MUTED)
    footer(s, 1, total)

    # 2 Agenda
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Agenda")
    items = [
        "Business problem",
        "Our solution",
        "Solution framework",
        "High-level architecture",
        "How questions become insights",
        "Question flow (DFD)",
        "Implementation approach",
        "Key features",
        "Product components (what we shipped)",
        "Legacy Streamlit → Ask DB",
        "Live demo / use cases",
        "Technology stack",
        "Future roadmap",
        "Ask / next steps",
    ]
    left_items = items[:7]
    right_items = items[7:]
    for i, t in enumerate(left_items):
        y = Inches(1.1) + Inches(i * 0.7)
        add_rect(s, Inches(0.7), y, Inches(0.45), Inches(0.45), ASK if i % 2 == 0 else DB)
        add_textbox(s, Inches(0.78), y + Inches(0.08), Inches(0.35), Inches(0.3), f"{i+1:02d}", size=11, bold=True, color=CARD, align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(1.35), y + Inches(0.08), Inches(4.5), Inches(0.35), t, size=16, color=INK)
    for i, t in enumerate(right_items):
        y = Inches(1.1) + Inches(i * 0.7)
        n = i + 8
        add_rect(s, Inches(7.0), y, Inches(0.45), Inches(0.45), DB if i % 2 == 0 else ASK)
        add_textbox(s, Inches(7.08), y + Inches(0.08), Inches(0.35), Inches(0.3), f"{n:02d}", size=11, bold=True, color=CARD, align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(7.65), y + Inches(0.08), Inches(4.8), Inches(0.35), t, size=16, color=INK)
    footer(s, 2, total)

    # 3 Business problem
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Business problem")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.35), "Today’s analytics loop is slow, brittle, and role-fragmented", size=14, color=MUTED)

    # Left cycle
    add_rect(s, Inches(0.5), Inches(1.35), Inches(5.8), Inches(5.2), CARD, LINE)
    add_textbox(s, Inches(0.75), Inches(1.5), Inches(5), Inches(0.35), "Current workflow", size=16, bold=True, color=ASK)
    steps = ["Business question", "Analyst / DE queue", "Hand-written SQL", "Dashboard refresh", "Answer (often late)"]
    for i, step in enumerate(steps):
        y = Inches(2.1) + Inches(i * 0.7)
        add_rect(s, Inches(1.1), y, Inches(4.4), Inches(0.5), RGBColor(0xEAF, 0xF2, 0xFF) if False else RGBColor(0xEA, 0xF2, 0xFF), LINE)
        add_textbox(s, Inches(1.25), y + Inches(0.1), Inches(4.1), Inches(0.35), f"{i+1}.  {step}", size=13, color=INK)
    add_textbox(s, Inches(1.1), Inches(5.8), Inches(4.5), Inches(0.4), "Turnaround often days → weeks", size=12, bold=True, color=WARN)

    # Right problems
    add_rect(s, Inches(6.6), Inches(1.35), Inches(6.2), Inches(5.2), CARD, LINE)
    add_textbox(s, Inches(6.85), Inches(1.5), Inches(5.5), Inches(0.35), "Friction between leadership & data teams", size=16, bold=True, color=DB)
    lead = [
        "“I don’t know SQL — I just need the number.”",
        "“When will I get my answer?”",
        "“Are there multiple sources of truth?”",
    ]
    data = [
        "Ambiguous business terms (FTR, loss ratio…)",
        "Same request repeated across analysts",
        "Too many ad-hoc asks to manage safely",
    ]
    add_textbox(s, Inches(6.85), Inches(2.1), Inches(5.5), Inches(0.3), "Leadership", size=12, bold=True, color=ASK)
    for i, t in enumerate(lead):
        add_textbox(s, Inches(6.85), Inches(2.45) + Inches(i * 0.45), Inches(5.6), Inches(0.4), f"• {t}", size=12, color=MUTED)
    add_textbox(s, Inches(6.85), Inches(4.0), Inches(5.5), Inches(0.3), "Analyst / Data Engineering", size=12, bold=True, color=DB)
    for i, t in enumerate(data):
        add_textbox(s, Inches(6.85), Inches(4.35) + Inches(i * 0.45), Inches(5.6), Inches(0.4), f"• {t}", size=12, color=MUTED)
    footer(s, 3, total)

    # 4 Our solution
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Our solution")
    add_textbox(
        s,
        Inches(0.5),
        Inches(0.85),
        Inches(12),
        Inches(0.35),
        "Enable natural-language access to enterprise data — with governance, not guesswork",
        size=14,
        color=MUTED,
    )
    path = [
        ("1. Ask", "Plain-English question in Chat / Home"),
        ("2. Understand", "Intent · entities · filters · glossary"),
        ("3. Plan", "Template-first or LLM SQL plan"),
        ("4. Generate SQL", "Schema-qualified, pack-grounded"),
        ("5. Validate", "Guardrails · EXPLAIN · auto-repair"),
        ("6. Execute", "Read-only PostgreSQL analytics DBs"),
        ("7. Present", "Table · Chart · SQL · Trust signals"),
        ("8. Act", "EI KPIs · Trust Center · Semantic Atlas"),
    ]
    for i, (title, desc) in enumerate(path):
        col = i % 4
        row = i // 4
        left = Inches(0.5) + Inches(col * 3.15)
        top = Inches(1.5) + Inches(row * 2.4)
        add_rect(s, left, top, Inches(3.0), Inches(2.05), CARD, LINE)
        add_rect(s, left, top, Inches(3.0), Inches(0.12), ASK if row == 0 else DB)
        add_textbox(s, left + Inches(0.2), top + Inches(0.35), Inches(2.6), Inches(0.4), title, size=16, bold=True, color=INK)
        add_textbox(s, left + Inches(0.2), top + Inches(0.9), Inches(2.6), Inches(0.9), desc, size=12, color=MUTED)
    footer(s, 4, total)

    # 5 Solution framework
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Solution framework")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.3), "User Intent  →  Semantic Layer  →  SQL Engine  →  Insight", size=14, bold=True, color=INK)

    layers = [
        ("INPUT", ASK, ["Natural language chat", "Industry switch (Auto / Insurance)", "Connected data sources status"]),
        ("SEMANTIC", PURPLE, ["YAML semantic packs", "Business glossary + ontology", "Value dictionary (live distincts)", "SQL guardrails (SELECT-only)"]),
        ("EXECUTION", DB, ["PostgreSQL analytics (RO)", "Template-first + LLM fallback", "Retry / self-heal EXPLAIN", "Query cache + SSE progress"]),
        ("PRESENTATION", ORANGE, ["Table / Chart / SQL tabs", "Executive Intelligence KPIs", "Data Trust Center", "Semantic Atlas · LLM Observability"]),
    ]
    for i, (name, color, items) in enumerate(layers):
        left = Inches(0.4) + Inches(i * 3.2)
        add_rect(s, left, Inches(1.35), Inches(3.05), Inches(4.7), CARD, LINE)
        add_rect(s, left, Inches(1.35), Inches(3.05), Inches(0.55), color)
        add_textbox(s, left + Inches(0.15), Inches(1.45), Inches(2.7), Inches(0.35), name, size=14, bold=True, color=CARD, align=PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            add_textbox(s, left + Inches(0.2), Inches(2.15) + Inches(j * 0.7), Inches(2.65), Inches(0.6), f"• {item}", size=12, color=INK)
    add_textbox(
        s,
        Inches(0.5),
        Inches(6.25),
        Inches(12),
        Inches(0.4),
        "Tech: Next.js · FastAPI · PostgreSQL · YAML semantic packs · Capgemini LLM API · TanStack Query · React Flow",
        size=12,
        color=MUTED,
    )
    footer(s, 5, total)

    # 6 High-level architecture
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "High-level architecture")
    boxes = [
        (0.5, 3.0, 2.0, 1.0, "User", ORANGE),
        (3.0, 3.0, 2.4, 1.0, "Ask DB Web\n(Next.js)", ASK),
        (6.0, 3.0, 2.6, 1.0, "Query Brain\n(FastAPI NLQ)", DB),
        (9.3, 1.5, 3.3, 1.0, "Semantic Atlas\n(YAML packs)", PURPLE),
        (9.3, 3.0, 3.3, 1.0, "PostgreSQL\nAnalytics (RO)", ORANGE),
        (9.3, 4.5, 3.3, 1.0, "App DB\nAuth · History · Cost", ASK),
        (6.0, 4.8, 2.6, 1.0, "Result UI\nTable · Chart · SQL", DB),
    ]
    for left, top, w, h, text, color in boxes:
        shape = add_rect(s, Inches(left), Inches(top), Inches(w), Inches(h), CARD, color)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = INK
        p.font.name = "Calibri"
        p.alignment = PP_ALIGN.CENTER
    notes = [
        "Browser never holds LLM keys or executes SQL",
        "Numbers always come from PostgreSQL — not the model",
        "Semantic truth from YAML packs, not ad-hoc prompts",
        "Analytics pools are read-only with statement timeout + LIMIT",
    ]
    for i, n in enumerate(notes):
        add_textbox(s, Inches(0.5), Inches(1.2) + Inches(i * 0.4), Inches(8.5), Inches(0.35), f"✓  {n}", size=13, color=MUTED)
    footer(s, 6, total)

    # 7 How questions become insights
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "How user questions become business insights")
    cols = [
        ("User", ASK, ["Natural language question", "Industry context", "Follow-ups"]),
        ("Intelligence", PURPLE, ["Question understanding", "Glossary + value match", "Semantic model / plan"]),
        ("Execution & Cache", DB, ["Cache hit?", "Template or LLM SQL", "Validate · repair · run"]),
        ("Data", ORANGE, ["Automotive PG", "Insurance PG", "Freshness / DQ signals"]),
        ("Visualization", WARN, ["Table (default)", "Interactive chart", "SQL + trust badges"]),
    ]
    for i, (title, color, items) in enumerate(cols):
        left = Inches(0.35) + Inches(i * 2.55)
        add_rect(s, left, Inches(1.2), Inches(2.4), Inches(5.3), CARD, LINE)
        add_rect(s, left, Inches(1.2), Inches(2.4), Inches(0.6), color)
        add_textbox(s, left + Inches(0.1), Inches(1.32), Inches(2.2), Inches(0.4), title, size=13, bold=True, color=CARD, align=PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            add_textbox(s, left + Inches(0.15), Inches(2.1) + Inches(j * 1.0), Inches(2.1), Inches(0.8), item, size=12, color=INK)
    footer(s, 7, total)

    # 8 Question flow DFD
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Question flow (DFD)")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.3), "From plain-English question to answered result in Chat", size=13, color=MUTED)

    flow_box(s, Inches(0.4), Inches(2.0), Inches(1.8), Inches(0.7), "User question", ASK, CARD)
    flow_box(s, Inches(2.6), Inches(2.0), Inches(1.6), Inches(0.7), "Router", DB, CARD)
    flow_box(s, Inches(4.7), Inches(1.1), Inches(2.0), Inches(0.65), "Polite redirect", ORANGE, CARD)
    flow_box(s, Inches(4.7), Inches(2.0), Inches(2.2), Inches(0.7), "Semantic layer", PURPLE, CARD)
    flow_box(s, Inches(4.7), Inches(3.1), Inches(2.2), Inches(0.7), "Edit prior SQL", DB, CARD)
    flow_box(s, Inches(7.4), Inches(2.0), Inches(2.4), Inches(0.7), "Generate / modify SQL", ASK, CARD)
    flow_box(s, Inches(10.2), Inches(2.0), Inches(2.5), Inches(0.7), "Postgres runs query", ORANGE, CARD)
    flow_box(s, Inches(10.2), Inches(3.4), Inches(2.5), Inches(0.7), "Table · Chart · SQL", DB, CARD)
    flow_box(s, Inches(7.4), Inches(4.6), Inches(5.3), Inches(0.7), "Answer in Chat (+ Trust / timing)", PURPLE, CARD)

    add_textbox(s, Inches(2.7), Inches(1.25), Inches(1.8), Inches(0.3), "off-topic", size=10, color=WARN)
    add_textbox(s, Inches(2.7), Inches(2.75), Inches(1.8), Inches(0.3), "new question", size=10, color=MUTED)
    add_textbox(s, Inches(2.7), Inches(3.85), Inches(2.0), Inches(0.3), "follow-up", size=10, color=MUTED)
    add_textbox(s, Inches(8.0), Inches(1.55), Inches(1.5), Inches(0.3), "safe SQL", size=10, color=OK)
    add_textbox(
        s,
        Inches(0.5),
        Inches(5.7),
        Inches(12),
        Inches(0.7),
        "Updated from legacy DuckDB path → PostgreSQL analytics with progressive SSE stages "
        "(understand → semantic match → generate → validate → crunch → build answer).",
        size=12,
        color=MUTED,
    )
    footer(s, 8, total)

    # 9 Implementation approach
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Implementation approach")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.3), "Data and AI interaction — production stack", size=13, color=MUTED)

    bands = [
        ("Presentation", ASK, "Next.js workspace · Home hero · Chat · EI · Trust · Semantic Atlas · LLM Observability"),
        ("Processing", RGBColor(0x3A, 0x4A, 0x5C), "FastAPI NLQ · templates · LLM · validation · SSE · cache · cost tracking"),
        ("AI", DB, "Capgemini LLM gateway · sampling controls · fallback model · usage metering"),
        ("Data", PURPLE, "askdb_app + askdb_automotive + askdb_insurance · Mongo/Qdrant for knowledge (optional)"),
    ]
    for i, (name, color, desc) in enumerate(bands):
        top = Inches(1.4) + Inches(i * 1.15)
        add_rect(s, Inches(0.5), top, Inches(12.3), Inches(1.0), CARD, LINE)
        add_rect(s, Inches(0.5), top, Inches(2.4), Inches(1.0), color)
        add_textbox(s, Inches(0.65), top + Inches(0.32), Inches(2.1), Inches(0.4), name, size=14, bold=True, color=CARD)
        add_textbox(s, Inches(3.15), top + Inches(0.3), Inches(9.3), Inches(0.5), desc, size=13, color=INK)
    footer(s, 9, total)

    # 10 Key features
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Key features")
    left_feats = [
        ("Intent recognition", "Understands ‘top performers’ as ORDER BY … DESC LIMIT — not keyword match."),
        ("Schema-aware context", "Pack columns, grains, and live value domains ground the LLM before it answers."),
        ("Governed relationships", "Join contracts + cardinality from Semantic Atlas — not silent cross joins."),
        ("Self-healing SQL", "Failed EXPLAIN / validation triggers repair; aborted transactions are recovered."),
        ("Config-driven domains", "Automotive & Insurance packs map business terms to physical columns."),
    ]
    right_feats = [
        ("Follow-up memory", "Prior SQL preserved so ‘only Mumbai’ revises the last successful query."),
        ("SQL guardrails", "SELECT-only · read-only txn · statement timeout · hard row LIMIT."),
        ("Executive Intelligence", "8 KPI cards, performance charts, India region → dealer → model map."),
        ("Data Trust Center", "Trust score, incidents, dataset health, profiling, lineage, rules."),
        ("Smart presentation", "Table-first results with Chart (X/Y/type) and SQL tabs; no fake confidence %."),
    ]
    for i, (t, d) in enumerate(left_feats):
        feature_row(s, Inches(0.5), Inches(1.15) + Inches(i * 1.05), Inches(5.7), t, d, ASK)
    for i, (t, d) in enumerate(right_feats):
        feature_row(s, Inches(6.8), Inches(1.15) + Inches(i * 1.05), Inches(5.9), t, d, DB)
    footer(s, 10, total)

    # 11 Product components
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Major product components")
    comps = [
        ("Home / Hero", "Ask DB branding, live demo loop, connected sources"),
        ("AI Chat", "SSE progress, Table/Chart/SQL, sticky layout"),
        ("Executive Intelligence", "KPIs, AI insights, performance analytics + map"),
        ("Data Trust Center", "Overview + Technical trust & DQ operations"),
        ("Semantic Atlas", "Ontology · Model · Glossary (cardinality)"),
        ("LLM Observability", "Cost analytics + sampling controls"),
        ("Data Preview", "Governed table browsing"),
        ("Activity", "Saved questions · history · system logs"),
    ]
    for i, (t, d) in enumerate(comps):
        col = i % 4
        row = i // 4
        left = Inches(0.45) + Inches(col * 3.2)
        top = Inches(1.3) + Inches(row * 2.55)
        add_rect(s, left, top, Inches(3.05), Inches(2.3), CARD, LINE)
        add_rect(s, left, top, Inches(3.05), Inches(0.1), ASK if col % 2 == 0 else DB)
        add_textbox(s, left + Inches(0.2), top + Inches(0.35), Inches(2.65), Inches(0.6), t, size=15, bold=True, color=INK)
        add_textbox(s, left + Inches(0.2), top + Inches(1.1), Inches(2.65), Inches(0.9), d, size=12, color=MUTED)
    footer(s, 11, total)

    # 12 Legacy vs new
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "From Streamlit AskData → Ask DB")
    add_rect(s, Inches(0.5), Inches(1.3), Inches(5.9), Inches(5.2), CARD, LINE)
    add_rect(s, Inches(6.9), Inches(1.3), Inches(5.9), Inches(5.2), CARD, LINE)
    add_textbox(s, Inches(0.7), Inches(1.45), Inches(5.4), Inches(0.4), "Legacy (Streamlit era)", size=16, bold=True, color=WARN)
    add_textbox(s, Inches(7.1), Inches(1.45), Inches(5.4), Inches(0.4), "Ask DB (current)", size=16, bold=True, color=DB)
    old = [
        "Streamlit UI",
        "DuckDB in-process engine",
        "CSV / file-oriented demos",
        "Dark PPT-era framing",
        "Monolithic notebook-style flow",
        "Limited enterprise shell",
    ]
    new = [
        "Next.js + design-system UI",
        "PostgreSQL analytics (RO pools)",
        "Industry semantic YAML packs",
        "Light premium product UI",
        "Modular services + SSE NLQ",
        "RBAC, Trust, EI, Observability",
    ]
    for i, t in enumerate(old):
        add_textbox(s, Inches(0.85), Inches(2.1) + Inches(i * 0.6), Inches(5.2), Inches(0.45), f"• {t}", size=14, color=MUTED)
    for i, t in enumerate(new):
        add_textbox(s, Inches(7.25), Inches(2.1) + Inches(i * 0.6), Inches(5.2), Inches(0.45), f"• {t}", size=14, color=INK)
    footer(s, 12, total)

    # 13 Use cases / demo
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Live demo / use cases")
    cases = [
        ("Automotive NLQ", ["Top SUV dealers in Mumbai", "Revenue by month", "Lowest body-style performers"]),
        ("Insurance NLQ", ["Loss ratio by product", "Claims by status", "Renewal / severity views"]),
        ("Executive room", ["8 KPI grid", "Performance charts with axes", "Region → dealer → model map"]),
        ("Trust & governance", ["Trust score hero", "Incidents & dataset health", "No bare confidence %"]),
    ]
    for i, (title, bullets) in enumerate(cases):
        left = Inches(0.45) + Inches((i % 2) * 6.4)
        top = Inches(1.25) + Inches((i // 2) * 2.7)
        bullet_card(s, left, top, Inches(6.1), Inches(2.45), title, bullets, ASK if i % 2 == 0 else DB)
    footer(s, 13, total)

    # 14 Tech stack
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Technology stack")
    stack = [
        ("Frontend", "Next.js App Router · React · Tailwind · TanStack Query · Zustand · React Flow"),
        ("API", "FastAPI · Pydantic v2 · SQLAlchemy 2 · SSE streaming · Argon2 / JWT auth"),
        ("Data", "PostgreSQL (app + automotive + insurance) · read-only analytics pools"),
        ("Semantic", "YAML packs · glossary · ontology compiler · value dictionary"),
        ("AI", "Capgemini LLM API · model fallback · temperature / top-p / top-k controls"),
        ("Ops", "Docker · health/readiness · cost metering · system logs · RBAC"),
    ]
    for i, (t, d) in enumerate(stack):
        top = Inches(1.2) + Inches(i * 0.85)
        add_rect(s, Inches(0.5), top, Inches(12.3), Inches(0.72), CARD, LINE)
        add_rect(s, Inches(0.5), top, Inches(2.3), Inches(0.72), ASK if i % 2 == 0 else DB)
        add_textbox(s, Inches(0.65), top + Inches(0.2), Inches(2.0), Inches(0.35), t, size=13, bold=True, color=CARD)
        add_textbox(s, Inches(3.0), top + Inches(0.2), Inches(9.5), Inches(0.4), d, size=13, color=INK)
    footer(s, 14, total)

    # 15 Roadmap
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    section_title(s, "Future roadmap")
    phases = [
        ("Now", "Stabilize NLQ accuracy · EI map · Trust · Ask DB brand polish"),
        ("Next", "Deeper access control on region/dealer drill-downs · richer chart packs"),
        ("Later", "What-If return · forecasting (optional) · durable metrics backend"),
        ("Scale", "Multi-tenant packs · warehouse connectors · production LLM budgets"),
    ]
    for i, (t, d) in enumerate(phases):
        left = Inches(0.45) + Inches(i * 3.2)
        add_rect(s, left, Inches(1.6), Inches(3.05), Inches(4.4), CARD, LINE)
        add_rect(s, left, Inches(1.6), Inches(3.05), Inches(0.7), ASK if i < 2 else DB)
        add_textbox(s, left + Inches(0.2), Inches(1.75), Inches(2.65), Inches(0.4), t, size=18, bold=True, color=CARD, align=PP_ALIGN.CENTER)
        add_textbox(s, left + Inches(0.25), Inches(2.7), Inches(2.55), Inches(2.8), d, size=14, color=INK)
    footer(s, 15, total)

    # 16 Close
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BG)
    add_rect(s, Inches(0), Inches(0), Inches(0.18), Inches(7.5), ASK)
    add_rect(s, Inches(0.18), Inches(0), Inches(0.08), Inches(7.5), DB)
    brand_title(s, Inches(0.8), Inches(2.3), Inches(11), Inches(0.8), prefix="", size=44)
    add_textbox(s, Inches(0.8), Inches(3.3), Inches(11), Inches(0.5), "Ready for leadership walkthrough & live demo", size=20, color=MUTED)
    add_textbox(
        s,
        Inches(0.8),
        Inches(4.2),
        Inches(11),
        Inches(1.0),
        "Ask a question in Chat · review Executive Intelligence · open Data Trust · browse Semantic Atlas",
        size=14,
        color=INK,
    )
    footer(s, 16, total)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
